"""Smart import: read any tabular export, find the real table, map its columns, clean its values.

The pipeline is:
    detect_file(path)  -> list[TableSpec]   (cheap: reads only the first rows of each sheet)
    iter_records(path, spec, quality)       (streams cleaned records for one table)

A TableSpec is plain JSON so it can be shown to the user, edited on the review page,
saved as a reusable mapping profile, and applied later.
"""
from __future__ import annotations

import calendar
import codecs
import csv
import hashlib
import json
import math
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

try:
    from python_calamine import CalamineWorkbook
except ImportError:  # pragma: no cover - calamine is in requirements, openpyxl is the fallback
    CalamineWorkbook = None

try:
    from rapidfuzz import fuzz

    def similarity(a: str, b: str) -> float:
        return fuzz.ratio(a, b) / 100
except ImportError:  # pragma: no cover
    from difflib import SequenceMatcher

    def similarity(a: str, b: str) -> float:
        return SequenceMatcher(None, a, b).ratio()


EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xlsb", ".xls", ".ods"}
TEXT_EXTENSIONS = {".csv", ".tsv", ".txt"}
JSON_EXTENSIONS = {".json"}
SUPPORTED_EXTENSIONS = EXCEL_EXTENSIONS | TEXT_EXTENSIONS | JSON_EXTENSIONS

HEADER_SCAN_ROWS = 40
PROFILE_SAMPLE_ROWS = 400
PREVIEW_ROWS = 6
ISSUE_SAMPLES = 5

# ---------------------------------------------------------------------------
# Canonical fields
# ---------------------------------------------------------------------------

FIELD_LABELS = {
    "ref": ("كود الصنف", "Item code"),
    "barcode": ("الباركود", "Barcode"),
    "name": ("اسم الصنف", "Item name"),
    "category": ("المجموعة / الفئة", "Category"),
    "supplier": ("المورد / الشركة", "Supplier"),
    "branch": ("الفرع", "Branch"),
    "qty_sold": ("الكمية المباعة", "Quantity sold"),
    "sales_value": ("قيمة المبيعات", "Sales value"),
    "customer_returns": ("مرتجع العملاء", "Customer returns"),
    "purchases": ("المشتريات", "Purchases"),
    "qoh": ("الرصيد الحالي", "Stock on hand"),
    "cost": ("تكلفة الوحدة", "Unit cost"),
    "price": ("سعر البيع", "Unit sale price"),
    "stock_value": ("قيمة المخزون بسعر البيع", "Stock value (sale price)"),
    "stock_cost_value": ("قيمة المخزون بالتكلفة", "Stock value (cost)"),
    "expiry": ("تاريخ الصلاحية", "Expiry date"),
    "batch": ("رقم التشغيلة", "Batch / lot"),
    "date": ("تاريخ الحركة", "Transaction date"),
    "invoice": ("رقم الفاتورة", "Invoice number"),
    "qty_generic": ("كمية (تتحدد تلقائي)", "Quantity (auto)"),
    "value_generic": ("قيمة (تتحدد تلقائي)", "Value (auto)"),
    "_ignore": ("تجاهل", "Ignore"),
}

NUMERIC_FIELDS = {
    "qty_sold", "sales_value", "customer_returns", "purchases", "qoh", "cost", "price",
    "stock_value", "stock_cost_value", "qty_generic", "value_generic",
}
DATE_FIELDS = {"expiry", "date"}
CODE_FIELDS = {"ref", "barcode", "invoice", "batch"}
TEXT_FIELDS = {"name", "category", "supplier", "branch"}
IDENTITY_FIELDS = ("ref", "barcode", "name")

SYNONYMS: dict[str, list[str]] = {
    "ref": [
        "ref", "reference", "product ref", "code", "item code", "product code", "item no", "item number",
        "item id", "product id", "sku", "stock code", "article", "article no", "plu", "internal code",
        "كود", "الكود", "كود الصنف", "رقم الصنف", "كود المنتج", "رمز الصنف", "الرمز", "كود الدواء",
    ],
    "barcode": ["barcode", "bar code", "ean", "ean13", "upc", "gtin", "international code", "الباركود", "باركود", "الكود الدولي"],
    "name": [
        "name", "item name", "product name", "item", "product", "description", "item description",
        "product description", "drug", "drug name", "medicine", "trade name",
        "الصنف", "اسم الصنف", "اسم المنتج", "المنتج", "البيان", "الاسم", "اسم الدواء", "الدواء", "الوصف", "الاسم التجاري",
    ],
    "category": [
        "category", "category1", "group", "item group", "class", "department", "family", "type", "section",
        "المجموعة", "مجموعة الصنف", "الفئة", "التصنيف", "القسم", "النوع", "المجموعه",
    ],
    "supplier": ["supplier", "vendor", "manufacturer", "company", "brand", "المورد", "الشركة", "الشركة المصنعة", "المصنع", "الشركه"],
    "branch": [
        "branch", "store", "location", "outlet", "shop", "site", "warehouse", "pharmacy", "branch name", "store name",
        "الفرع", "اسم الفرع", "المخزن", "الصيدلية", "الموقع", "المستودع",
    ],
    "qty_sold": [
        "sold", "qty sold", "quantity sold", "units sold", "sales qty", "sales quantity", "pos sales", "sold qty",
        "sale qty", "sales units", "sales", "total sold",
        "الكمية المباعة", "المباع", "كمية المبيعات", "مبيعات", "عدد المباع", "الكميه المباعه", "كمية البيع", "المبيعات",
    ],
    "sales_value": [
        "net sales", "sales value", "sales amount", "revenue", "total sales", "sales total", "turnover", "gross sales",
        "قيمة المبيعات", "صافي المبيعات", "اجمالي المبيعات", "اجمالي البيع", "الايراد", "الايرادات", "قيمة البيع",
    ],
    "customer_returns": ["customer returns", "returns", "return qty", "sales returns", "مرتجع", "المرتجعات", "مرتجع العملاء", "مرتجع مبيعات"],
    "purchases": ["purchases", "purchased qty", "received qty", "purchase qty", "المشتريات", "الوارد", "كمية المشتريات"],
    "qoh": [
        "stock", "on hand", "qoh", "qty on hand", "quantity on hand", "balance", "stock qty", "current stock",
        "available", "available qty", "inventory", "closing stock", "stock on hand", "soh", "end qty",
        "الرصيد", "الرصيد الحالي", "رصيد", "المخزون", "الكمية المتاحة", "الموجود", "الكمية الحالية", "رصيد المخزن",
        "رصيد اخر المدة", "الرصيد الفعلي", "المتاح",
    ],
    "cost": [
        "cost", "unit cost", "cost price", "avg cost", "average cost", "purchase price", "buying price", "buy price",
        "wac", "cost per unit",
        "سعر الشراء", "التكلفة", "تكلفة الوحدة", "سعر التكلفة", "متوسط التكلفة", "التكلفه", "سعر الصيدلي", "سعر الصيدلية",
    ],
    "price": [
        "price", "sales price", "sale price", "selling price", "retail price", "unit price", "retail", "public price",
        "rsp", "mrp", "sell price",
        "سعر البيع", "السعر", "سعر الوحدة", "سعر الجمهور", "سعر التجزئة", "سعر المستهلك",
    ],
    "stock_value": [
        "total sales price", "stock value", "sale price balance", "inventory value", "stock sales value", "retail value",
        "قيمة المخزون", "قيمة الرصيد", "قيمة الرصيد بسعر البيع",
    ],
    "stock_cost_value": [
        "cost price balance", "stock cost", "cost value", "total cost", "inventory cost", "stock cost value",
        "قيمة التكلفة", "تكلفة المخزون", "اجمالي التكلفة", "قيمة الرصيد بالتكلفة",
    ],
    "expiry": [
        "expiry", "expiry date", "exp", "exp date", "expiration", "expiration date", "expire", "expires", "best before",
        "use by", "صلاحية", "الصلاحية", "تاريخ الصلاحية", "تاريخ الانتهاء", "تاريخ انتهاء الصلاحية", "انتهاء", "ينتهي",
    ],
    "batch": ["batch", "batch no", "lot", "lot no", "lot number", "batch number", "التشغيلة", "رقم التشغيلة", "الباتش", "اللوت"],
    "date": [
        "date", "invoice date", "transaction date", "sale date", "sales date", "day", "created at", "trans date", "doc date",
        "التاريخ", "تاريخ الفاتورة", "تاريخ البيع", "اليوم", "تاريخ الحركة",
    ],
    "invoice": [
        "invoice", "invoice no", "invoice number", "receipt", "receipt no", "bill", "bill no", "order id",
        "transaction id", "doc no", "رقم الفاتورة", "الفاتورة", "رقم الايصال", "فاتورة", "رقم الحركة",
    ],
    "qty_generic": ["qty", "quantity", "units", "pcs", "count", "الكمية", "كمية", "العدد", "الكميه"],
    "value_generic": ["value", "amount", "total", "line total", "الاجمالي", "القيمة", "المبلغ", "الاجمالى", "القيمه"],
    "_ignore": [
        "م", "مسلسل", "serial", "no", "#", "sr", "rank", "customer id", "customer", "country", "notes", "remarks",
        "ملاحظات", "supplier returns", "adjustments", "internal receives", "internal sends", "b2b sales", "tax",
        "vat", "الضريبة", "العميل", "discount", "الخصم", "unit", "الوحدة",
    ],
}

TOTAL_WORDS = {
    "total", "totals", "grand total", "subtotal", "sub total", "sum",
    "الاجمالي", "اجمالي", "المجموع", "مجموع", "الاجمالي العام", "اجمالي عام", "الاجمالى",
}

# ---------------------------------------------------------------------------
# Text normalization
# ---------------------------------------------------------------------------

ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬", "01234567890123456789.,")
ARABIC_LETTERS = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ة": "ه", "ى": "ي", "ؤ": "و", "ئ": "ي", "ـ": ""})
TASHKEEL = re.compile(r"[ؐ-ًؚ-ٰٟۖ-ۭ]")
CAMEL = re.compile(r"(?<=[a-z])(?=[A-Z])")
NON_WORD = re.compile(r"[^\w#]+", flags=re.UNICODE)


def normalize_text(value: Any) -> str:
    text = str(value if value is not None else "")
    text = CAMEL.sub(" ", text)
    text = TASHKEEL.sub("", text).translate(ARABIC_LETTERS).translate(ARABIC_DIGITS).casefold()
    text = text.replace("_", " ")
    text = NON_WORD.sub(" ", text)
    return re.sub(r"\s+", " ", text).strip()


NORMALIZED_SYNONYMS = {field: sorted({normalize_text(s) for s in values}, key=len, reverse=True) for field, values in SYNONYMS.items()}
NORMALIZED_TOTAL_WORDS = {normalize_text(word) for word in TOTAL_WORDS}


def strip_definite_article(token: str) -> str:
    return token[2:] if token.startswith("ال") and len(token) > 3 else token


def header_field_scores(header: Any) -> dict[str, float]:
    """Score how well one header cell matches each canonical field (0..1)."""
    norm = normalize_text(header)
    if not norm:
        return {}
    tokens = {strip_definite_article(token) for token in norm.split()}
    compact = norm.replace(" ", "")
    scores: dict[str, float] = {}
    for field_name, synonyms in NORMALIZED_SYNONYMS.items():
        best = 0.0
        for synonym in synonyms:
            if norm == synonym or compact == synonym.replace(" ", ""):
                best = 1.0
                break
            syn_tokens = {strip_definite_article(token) for token in synonym.split()}
            if syn_tokens and syn_tokens <= tokens:
                coverage = len(syn_tokens) / max(len(tokens), 1)
                # single generic tokens inside a long header are weak evidence
                best = max(best, 0.62 + 0.3 * coverage)
                continue
            if len(norm) >= 4 and len(synonym) >= 4:
                ratio = similarity(norm, synonym)
                if ratio >= 0.86:
                    best = max(best, ratio * 0.85)
        if best > 0:
            scores[field_name] = round(best, 3)
    return scores


# ---------------------------------------------------------------------------
# Value cleaning
# ---------------------------------------------------------------------------

NUMBER_STRIP = re.compile(r"[^0-9.,]")
SCIENTIFIC = re.compile(r"^\s*-?\d+(?:[.,]\d+)?[eE][+-]?\d+\s*$")
THOUSANDS_COMMA = re.compile(r"^\d{1,3}(,\d{3})+(\.\d+)?$")
EXCEL_EPOCH = date(1899, 12, 30)


def clean_number(value: Any) -> float | None:
    """Parse numbers like '1,250.50', '١٬٢٥٠٫٥', '(300)', '45 ج.م', '1.234,56', '12-'. None if not a number."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return None if isinstance(value, float) and (math.isnan(value) or math.isinf(value)) else float(value)
    if isinstance(value, (date, datetime)):
        return None
    text = str(value).translate(ARABIC_DIGITS).strip()
    if not text:
        return None
    if SCIENTIFIC.match(text):
        try:
            return float(text.replace(",", "."))
        except ValueError:
            return None
    negative = (text.startswith("(") and text.endswith(")")) or text.startswith("-") or text.endswith("-") or text.startswith("−")
    digits = NUMBER_STRIP.sub("", text).strip(".,")
    if not re.search(r"\d", digits):
        return None
    if "," in digits and "." in digits:
        if digits.rfind(",") > digits.rfind("."):
            digits = digits.replace(".", "").replace(",", ".")
        else:
            digits = digits.replace(",", "")
    elif "," in digits:
        if THOUSANDS_COMMA.match(digits):
            digits = digits.replace(",", "")
        elif digits.count(",") == 1 and len(digits.split(",")[1]) <= 2:
            digits = digits.replace(",", ".")
        else:
            digits = digits.replace(",", "")
    elif digits.count(".") > 1:
        digits = digits.replace(".", "")
    try:
        number = float(digits)
    except ValueError:
        return None
    return -number if negative else number


DATE_FORMATS_DAYFIRST = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y"]
DATE_FORMATS_MONTHFIRST = ["%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y"]
DATE_FORMATS_ISO = ["%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%d %b %Y", "%d %B %Y", "%b %d %Y", "%d-%b-%Y", "%d-%b-%y"]
MONTH_FORMATS = ["%m/%Y", "%m-%Y", "%Y-%m", "%Y/%m", "%m/%y", "%m-%y", "%b %Y", "%B %Y", "%b-%Y", "%b-%y", "%b %y"]
DATE_SPLIT = re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})")


def clean_date(value: Any, dayfirst: bool = True, month_end: bool = False) -> date | None:
    """Parse dates from Excel cells, serial numbers, and many text formats (incl. month/year)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20000 <= value <= 80000:
            return EXCEL_EPOCH + timedelta(days=int(value))
        if 19000101 <= value <= 21001231:
            try:
                return datetime.strptime(str(int(value)), "%Y%m%d").date()
            except ValueError:
                return None
        return None
    text = str(value).translate(ARABIC_DIGITS).strip()
    if not text:
        return None
    text = re.sub(r"[T\s]+\d{1,2}:\d{2}(:\d{2})?(\.\d+)?\s*(am|pm|ص|م)?\s*$", "", text, flags=re.IGNORECASE).strip()
    text = re.sub(r"\s+", " ", text)
    formats = (DATE_FORMATS_DAYFIRST + DATE_FORMATS_MONTHFIRST) if dayfirst else (DATE_FORMATS_MONTHFIRST + DATE_FORMATS_DAYFIRST)
    for fmt in DATE_FORMATS_ISO + formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    for fmt in MONTH_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        if month_end:
            return parsed.replace(day=calendar.monthrange(parsed.year, parsed.month)[1])
        return parsed
    return None


def clean_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    text = " ".join(str(value).translate(ARABIC_DIGITS).split())
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def clean_label(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------

def detect_encoding(path: Path) -> str:
    raw = path.read_bytes()[:200_000]
    if raw.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if raw.startswith(codecs.BOM_UTF16_LE) or raw.startswith(codecs.BOM_UTF16_BE):
        return "utf-16"
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError as exc:
        # A multi-byte character cut at the sample boundary is still UTF-8.
        if exc.start >= len(raw) - 4:
            return "utf-8"
    try:
        raw.decode("cp1256")
        return "cp1256"
    except UnicodeDecodeError:
        return "latin-1"


def detect_delimiter(sample: str, extension: str) -> str:
    if extension == ".tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        counts = {delim: sample.count(delim) for delim in ",;\t|"}
        return max(counts, key=counts.get)


def sheet_readers(path: Path) -> list[tuple[str, Callable[[], Iterator[list[Any]]]]]:
    """Return (sheet name, row-iterator factory) for every sheet/table in the file."""
    extension = path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        encoding = detect_encoding(path)
        with path.open(encoding=encoding, errors="replace", newline="") as handle:
            sample = handle.read(20_000)
        delimiter = detect_delimiter(sample, extension)

        def csv_rows() -> Iterator[list[Any]]:
            with path.open(encoding=encoding, errors="replace", newline="") as handle:
                for row in csv.reader(handle, delimiter=delimiter):
                    yield row

        return [(path.name, csv_rows)]

    if extension in JSON_EXTENSIONS:
        data = json.loads(path.read_text(encoding=detect_encoding(path)))
        if isinstance(data, dict):
            data = next((value for value in data.values() if isinstance(value, list)), [data])
        records = [row for row in data if isinstance(row, dict)]
        headers: list[str] = []
        for record in records:
            for key in record:
                if key not in headers:
                    headers.append(key)

        def json_rows() -> Iterator[list[Any]]:
            yield list(headers)
            for record in records:
                yield [record.get(key) for key in headers]

        return [(path.name, json_rows)]

    if extension in EXCEL_EXTENSIONS:
        if CalamineWorkbook is not None:
            try:
                workbook = CalamineWorkbook.from_path(str(path))
                readers = []
                for sheet_name in workbook.sheet_names:
                    def calamine_rows(name: str = sheet_name) -> Iterator[list[Any]]:
                        book = CalamineWorkbook.from_path(str(path))
                        try:
                            yield from book.get_sheet_by_name(name).iter_rows()
                        finally:
                            book.close()
                    readers.append((sheet_name, calamine_rows))
                workbook.close()
                return readers
            except Exception:
                if extension not in {".xlsx", ".xlsm"}:
                    raise
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        names = workbook.sheetnames
        workbook.close()
        readers = []
        for sheet_name in names:
            def openpyxl_rows(name: str = sheet_name) -> Iterator[list[Any]]:
                book = load_workbook(path, read_only=True, data_only=True)
                try:
                    for row in book[name].iter_rows(values_only=True):
                        yield list(row)
                finally:
                    book.close()
            readers.append((sheet_name, openpyxl_rows))
        return readers

    raise ValueError(f"Unsupported file type: {extension or path.name}")


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


# ---------------------------------------------------------------------------
# Detection
# ---------------------------------------------------------------------------

@dataclass
class ColumnSpec:
    index: int
    header: str
    field: str = ""          # canonical field or "" (unmapped)
    confidence: float = 0.0
    method: str = ""         # header | content | profile | user
    dayfirst: bool = True
    kind: str = ""           # number | date | text | code | empty


@dataclass
class TableSpec:
    file_name: str
    sheet: str
    header_row: int                      # 0-based index of the header row (-1 = no header)
    columns: list[ColumnSpec]
    usable: bool = True
    reason: str = ""
    role: list[str] = field(default_factory=list)
    meta: dict[str, str] = field(default_factory=dict)
    preview: list[dict[str, Any]] = field(default_factory=list)
    signature: str = ""
    profile_applied: bool = False

    def mapping(self) -> dict[str, int]:
        return {col.field: col.index for col in self.columns if col.field and col.field != "_ignore"}

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "TableSpec":
        columns = [ColumnSpec(**column) for column in data.get("columns", [])]
        return cls(**{**data, "columns": columns})


def header_signature(headers: list[str]) -> str:
    normalized = "|".join(normalize_text(header) for header in headers)
    return hashlib.sha1(normalized.encode("utf-8")).hexdigest()[:16]


def row_header_score(row: list[Any]) -> tuple[float, int]:
    strong = 0
    total = 0.0
    text_cells = 0
    for cell in row:
        if is_blank(cell):
            continue
        if isinstance(cell, (int, float, date, datetime)) and not isinstance(cell, bool):
            total -= 0.3
            continue
        text_cells += 1
        scores = header_field_scores(cell)
        if not scores:
            continue
        best = max(scores.values())
        if best >= 0.8:
            strong += 1
        total += best
    return total, strong


def find_header_row(rows: list[list[Any]]) -> int:
    best_index, best_score, best_strong = -1, 0.0, 0
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score, strong = row_header_score(row)
        if strong >= 2 and score > best_score:
            best_index, best_score, best_strong = index, score, strong
    return best_index


def profile_column(values: list[Any]) -> dict[str, Any]:
    non_empty = [value for value in values if not is_blank(value)]
    if not non_empty:
        return {"kind": "empty", "fill": 0.0}
    numbers = dates = texts = 0
    digit_lengths: Counter[int] = Counter()
    first_gt_12 = second_gt_12 = 0
    future = 0
    today = date.today()
    text_lengths = []
    for value in non_empty:
        if isinstance(value, (datetime, date)):
            dates += 1
            if (value.date() if isinstance(value, datetime) else value) > today:
                future += 1
            continue
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            numbers += 1
            if float(value).is_integer() and value >= 0:
                digit_lengths[len(str(int(value)))] += 1
            continue
        text = str(value).translate(ARABIC_DIGITS).strip()
        match = DATE_SPLIT.match(text)
        parsed = clean_date(text, month_end=True) if (match or re.match(r"^\d{4}[-/.]\d{1,2}", text) or re.match(r"^\d{1,2}[/-]\d{2,4}$", text)) else None
        if parsed:
            dates += 1
            if parsed > today:
                future += 1
            if match:
                if int(match.group(1)) > 12:
                    first_gt_12 += 1
                if int(match.group(2)) > 12:
                    second_gt_12 += 1
            continue
        number = clean_number(text) if re.search(r"\d", text) and not re.search(r"[a-zA-Zء-ي]{3,}", text) else None
        if number is not None:
            numbers += 1
            if re.fullmatch(r"\d+", text):
                digit_lengths[len(text)] += 1
            continue
        texts += 1
        text_lengths.append(len(text))
    count = len(non_empty)
    kind = "text"
    if dates / count >= 0.7:
        kind = "date"
    elif numbers / count >= 0.7:
        kind = "number"
    common_length, common_count = digit_lengths.most_common(1)[0] if digit_lengths else (0, 0)
    if kind == "number" and 8 <= common_length <= 14 and common_count / count >= 0.7:
        kind = "code"
    return {
        "kind": kind,
        "fill": count / max(len(values), 1),
        "number_share": numbers / count,
        "date_share": dates / count,
        "text_share": texts / count,
        "future_share": future / max(dates, 1),
        "avg_text_len": sum(text_lengths) / len(text_lengths) if text_lengths else 0,
        "dayfirst": not (second_gt_12 > 0 and first_gt_12 == 0),
        "distinct": len({str(value) for value in non_empty}),
    }


def assign_fields(headers: list[str], profiles: list[dict[str, Any]]) -> list[ColumnSpec]:
    columns = [
        ColumnSpec(index=index, header=clean_label(header), kind=profiles[index]["kind"], dayfirst=profiles[index].get("dayfirst", True))
        for index, header in enumerate(headers)
    ]
    candidates: list[tuple[float, int, str]] = []
    for index, header in enumerate(headers):
        for field_name, score in header_field_scores(header).items():
            profile = profiles[index]
            if profile["kind"] == "empty" and field_name != "_ignore":
                score -= 0.3
            if field_name in NUMERIC_FIELDS and profile["kind"] not in {"number", "code", "empty"}:
                score -= 0.35
            if field_name in DATE_FIELDS and profile["kind"] not in {"date", "empty"}:
                score -= 0.35
            if field_name == "name" and profile["kind"] in {"number", "code", "date"}:
                score -= 0.4
            if score >= 0.6:
                candidates.append((score, index, field_name))
    taken_fields: set[str] = set()
    for score, index, field_name in sorted(candidates, key=lambda c: (-c[0], c[1])):
        column = columns[index]
        if column.field or (field_name in taken_fields and field_name != "_ignore"):
            continue
        column.field, column.confidence, column.method = field_name, round(min(score, 1.0), 2), "header"
        taken_fields.add(field_name)

    # Content inference for columns the headers did not explain.
    def free(field_name: str) -> bool:
        return field_name not in taken_fields

    unmapped = [column for column in columns if not column.field and column.kind != "empty"]
    for column in unmapped:
        profile = profiles[column.index]
        if column.kind == "code" and free("barcode"):
            column.field, column.confidence, column.method = "barcode", 0.6, "content"
        elif column.kind == "date":
            if profile["future_share"] >= 0.5 and free("expiry"):
                column.field, column.confidence, column.method = "expiry", 0.6, "content"
            elif free("date"):
                column.field, column.confidence, column.method = "date", 0.55, "content"
        if column.field:
            taken_fields.add(column.field)
    if free("name"):
        text_columns = [c for c in columns if not c.field and c.kind == "text" and profiles[c.index]["avg_text_len"] >= 6]
        if text_columns:
            best = max(text_columns, key=lambda c: (profiles[c.index]["distinct"], profiles[c.index]["avg_text_len"]))
            best.field, best.confidence, best.method = "name", 0.5, "content"
    return columns


def resolve_generic_fields(columns: list[ColumnSpec]) -> None:
    """'Quantity' means sold units in an invoice file and on-hand units in a stock file."""
    present = {column.field for column in columns}
    transactional = "date" in present or "invoice" in present
    for column in columns:
        if column.field == "qty_generic":
            target = "qty_sold" if transactional else "qoh"
            if target in present:
                target = "qoh" if target == "qty_sold" and "qoh" not in present else ""
            column.field = target or "_ignore"
            present.add(column.field)
        elif column.field == "value_generic":
            target = "sales_value" if transactional or "qoh" not in present else "stock_value"
            if target in present:
                target = ""
            column.field = target or "_ignore"
            present.add(column.field)


def table_role(columns: list[ColumnSpec]) -> list[str]:
    present = {column.field for column in columns}
    role = []
    if ("date" in present or "invoice" in present) and ("qty_sold" in present or "sales_value" in present):
        role.append("transactions")
    elif "qty_sold" in present or "sales_value" in present:
        role.append("sales")
    if "qoh" in present:
        role.append("stock")
    if "expiry" in present:
        role.append("expiry")
    if "cost" in present or "stock_cost_value" in present:
        role.append("cost")
    if "price" in present:
        role.append("price")
    return role


def validate_spec(spec: TableSpec) -> None:
    present = {column.field for column in spec.columns}
    spec.role = table_role(spec.columns)
    if not any(field_name in present for field_name in IDENTITY_FIELDS):
        spec.usable, spec.reason = False, "no_identity"
    elif not spec.role:
        spec.usable, spec.reason = False, "no_measures"
    else:
        spec.usable, spec.reason = True, ""


PERIOD_IN_HEADER = re.compile(r"\b(\d{1,3})\s*(?:d|day|days|يوم|ايام)\b")
INTRO_DATE = re.compile(r"(\d{1,4}[/\-.]\d{1,2}[/\-.]\d{1,4})")


def extract_intro_meta(rows: list[list[Any]]) -> dict[str, str]:
    """Pick up report period and branch from title rows above the table."""
    found_dates: list[date] = []
    meta: dict[str, str] = {}
    for row_index, row in enumerate(rows):
        cells = [cell for cell in row if not is_blank(cell)]
        for position, cell in enumerate(row):
            if isinstance(cell, (datetime, date)):
                found_dates.append(cell.date() if isinstance(cell, datetime) else cell)
                continue
            if is_blank(cell):
                continue
            text = str(cell).translate(ARABIC_DIGITS)
            stamp_cell = re.search(r"report time|generated|printed|print date|تاريخ الطباعه|وقت التقرير", normalize_text(text))
            for match in ([] if stamp_cell else INTRO_DATE.findall(text)):
                parsed = clean_date(match)
                if parsed:
                    found_dates.append(parsed)
            norm = normalize_text(text)
            if norm in {"branch", "الفرع", "store"}:
                below = rows[row_index + 1] if row_index + 1 < len(rows) else []
                candidates = [row[position + 1] if position + 1 < len(row) else None, below[position] if position < len(below) else None]
                for candidate in candidates:
                    if not is_blank(candidate) and max(header_field_scores(candidate).values(), default=0) < 0.95:
                        meta.setdefault("branch_hint", clean_label(candidate))
                        break
            branch_match = re.search(r"(?:فرع|branch)\s*[:\-]?\s*([^\-|،,]+)$", text.strip(), flags=re.IGNORECASE)
            if branch_match and len(cells) <= 2:
                meta.setdefault("branch_hint", branch_match.group(1).strip())
    distinct = sorted(set(found_dates))
    if len(distinct) >= 2:
        meta["start_date"] = distinct[0].isoformat()
        meta["end_date"] = distinct[-1].isoformat()
    return meta


def detect_table(file_name: str, sheet: str, row_factory: Callable[[], Iterator[list[Any]]], profiles_store: dict[str, Any] | None = None) -> TableSpec:
    head: list[list[Any]] = []
    for row in row_factory():
        head.append(list(row))
        if len(head) >= HEADER_SCAN_ROWS + PROFILE_SAMPLE_ROWS:
            break
    header_index = find_header_row(head)
    if header_index < 0:
        width = max((len(row) for row in head), default=0)
        headers = [f"Column {index + 1}" for index in range(width)]
        body = head
        intro: list[list[Any]] = []
    else:
        headers = [clean_label(cell) for cell in head[header_index]]
        body = head[header_index + 1:]
        intro = head[:header_index]
    width = max([len(headers)] + [len(row) for row in body[:50]])
    headers += [f"Column {index + 1}" for index in range(len(headers), width)]
    body = [row for row in body if not all(is_blank(cell) for cell in row)]
    profiles = [profile_column([row[index] if index < len(row) else None for row in body]) for index in range(width)]

    spec = TableSpec(file_name=file_name, sheet=sheet, header_row=header_index, columns=[], signature=header_signature(headers))
    saved = (profiles_store or {}).get(spec.signature)
    if saved:
        spec.columns = [ColumnSpec(**column) for column in saved["columns"]]
        spec.profile_applied = True
    else:
        spec.columns = assign_fields(headers, profiles)
        resolve_generic_fields(spec.columns)
    spec.meta = extract_intro_meta(intro)
    for column in spec.columns:
        if column.field in {"qty_sold", "sales_value"}:
            match = PERIOD_IN_HEADER.search(normalize_text(column.header))
            if match:
                spec.meta["period_days"] = match.group(1)
    validate_spec(spec)
    spec.preview = build_preview(spec, body[:PREVIEW_ROWS * 3])
    return spec


def detect_file(path: Path, display_name: str | None = None, profiles_store: dict[str, Any] | None = None) -> list[TableSpec]:
    specs = []
    for sheet, factory in sheet_readers(path):
        spec = detect_table(display_name or path.name, sheet, factory, profiles_store)
        specs.append(spec)
    return specs


# ---------------------------------------------------------------------------
# Record streaming + data quality
# ---------------------------------------------------------------------------

@dataclass
class Quality:
    rows_read: int = 0
    rows_used: int = 0
    counters: Counter = field(default_factory=Counter)
    samples: dict[str, list[str]] = field(default_factory=dict)

    def flag(self, issue: str, sample: str | None = None) -> None:
        self.counters[issue] += 1
        if sample is not None:
            bucket = self.samples.setdefault(issue, [])
            if len(bucket) < ISSUE_SAMPLES:
                bucket.append(sample)

    def merge(self, other: "Quality") -> None:
        self.rows_read += other.rows_read
        self.rows_used += other.rows_used
        self.counters.update(other.counters)
        for issue, samples in other.samples.items():
            bucket = self.samples.setdefault(issue, [])
            bucket.extend(samples[: max(ISSUE_SAMPLES - len(bucket), 0)])

    def to_dict(self) -> dict[str, Any]:
        return {"rows_read": self.rows_read, "rows_used": self.rows_used, "counters": dict(self.counters), "samples": self.samples}


def convert_cell(field_name: str, value: Any, column: ColumnSpec) -> tuple[Any, bool]:
    """Return (clean value, ok). ok=False means a non-empty cell could not be understood."""
    if field_name in NUMERIC_FIELDS:
        if is_blank(value):
            return None, True
        number = clean_number(value)
        return number, number is not None
    if field_name in DATE_FIELDS:
        if is_blank(value):
            return None, True
        parsed = clean_date(value, dayfirst=column.dayfirst, month_end=field_name == "expiry")
        return parsed, parsed is not None
    if field_name in CODE_FIELDS:
        return clean_code(value), True
    return clean_label(value), True


def row_label(record: dict[str, Any], row_number: int) -> str:
    identity = record.get("name") or record.get("ref") or record.get("barcode") or ""
    return f"#{row_number}: {identity}".strip()


TOTAL_HINT = re.compile(r"total|sum|جمال|جمل|مجموع", flags=re.IGNORECASE)


def is_totals_row(texts: list[str]) -> bool:
    for text in texts:
        cell = normalize_text(text)
        if cell in NORMALIZED_TOTAL_WORDS or cell.startswith("total ") or cell.startswith("اجمالي "):
            return True
    return False


def iter_records(path: Path, spec: TableSpec, quality: Quality) -> Iterator[dict[str, Any]]:
    readers = dict(sheet_readers(path))
    factory = readers.get(spec.sheet)
    if factory is None:
        raise ValueError(f"Sheet not found: {spec.sheet}")
    columns = [column for column in spec.columns if column.field and column.field != "_ignore"]
    header_raw = [column.header.strip().casefold() for column in spec.columns]
    header_needed = max(2, len(header_raw) // 2)
    for row_index, row in enumerate(factory()):
        if row_index <= spec.header_row:
            continue
        row = list(row)
        row_number = row_index + 1
        if all(is_blank(cell) for cell in row):
            continue
        quality.rows_read += 1
        text_cells = [cell for cell in row if isinstance(cell, str) and cell.strip()]
        if spec.header_row >= 0 and len(text_cells) >= header_needed and sum(
            1 for a, b in zip(row, header_raw) if b and isinstance(a, str) and a.strip().casefold() == b
        ) >= header_needed:
            quality.flag("repeated_header", f"#{row_number}")
            continue
        if any(TOTAL_HINT.search(cell) for cell in text_cells[:4]) and is_totals_row(text_cells[:4]):
            ident = [row[c.index] for c in columns if c.field in ("ref", "barcode") and c.index < len(row)]
            if not any(not is_blank(cell) and normalize_text(cell) not in NORMALIZED_TOTAL_WORDS for cell in ident):
                quality.flag("totals_row", f"#{row_number}")
                continue
        record: dict[str, Any] = {"_row": row_number}
        bad_fields = []
        for column in columns:
            value = row[column.index] if column.index < len(row) else None
            clean, ok = convert_cell(column.field, value, column)
            if not ok:
                bad_fields.append(column.field)
                quality.flag(f"bad_{'date' if column.field in DATE_FIELDS else 'number'}", f"#{row_number} {column.header}={value!r}")
            record[column.field] = clean
        if not any(record.get(field_name) for field_name in IDENTITY_FIELDS):
            quality.flag("no_identity", f"#{row_number}")
            continue
        barcode = record.get("barcode") or ""
        if barcode and re.fullmatch(r"\d+(\.\d+)?[eE]\+?\d+", barcode):
            quality.flag("barcode_scientific", row_label(record, row_number))
            record["barcode"] = ""
        if (record.get("qoh") or 0) < 0:
            quality.flag("negative_stock", row_label(record, row_number))
        quality.rows_used += 1
        yield record


# ---------------------------------------------------------------------------
# Preview + mapping profiles
# ---------------------------------------------------------------------------

def build_preview(spec: TableSpec, body: list[list[Any]]) -> list[dict[str, Any]]:
    preview = []
    columns = [column for column in spec.columns if column.field and column.field != "_ignore"]
    for row in body:
        record = {}
        for column in columns:
            value = row[column.index] if column.index < len(row) else None
            clean, _ok = convert_cell(column.field, value, column)
            record[column.field] = clean.isoformat() if isinstance(clean, date) else clean
        if any(record.get(field_name) for field_name in IDENTITY_FIELDS):
            preview.append(record)
        if len(preview) >= PREVIEW_ROWS:
            break
    return preview


def apply_user_mapping(spec: TableSpec, overrides: dict[int, str]) -> TableSpec:
    for column in spec.columns:
        if column.index not in overrides:
            continue
        chosen = overrides[column.index]
        if chosen != column.field:
            column.field = chosen
            column.confidence = 1.0 if chosen else 0.0
            column.method = "user"
    validate_spec(spec)
    return spec


def load_profiles(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def save_profile(path: Path, spec: TableSpec) -> None:
    profiles = load_profiles(path)
    profiles[spec.signature] = {
        "file_name": spec.file_name,
        "sheet": spec.sheet,
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "columns": [asdict(column) for column in spec.columns],
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(profiles, ensure_ascii=False, indent=2), encoding="utf-8")


def field_choices(lang: str = "ar") -> list[tuple[str, str]]:
    position = 0 if lang == "ar" else 1
    return [("", "—")] + [(key, labels[position]) for key, labels in FIELD_LABELS.items() if key not in {"qty_generic", "value_generic"}]


def iter_all_records(paths_specs: Iterable[tuple[Path, TableSpec]], quality: Quality) -> Iterator[tuple[TableSpec, dict[str, Any]]]:
    for path, spec in paths_specs:
        for record in iter_records(path, spec, quality):
            yield spec, record
