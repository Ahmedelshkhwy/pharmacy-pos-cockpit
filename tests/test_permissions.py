"""Roles: a branch manager or viewer must never see or change another branch."""
from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "sample_data"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from test_web_flow import client, submit_review, upload  # noqa: E402,F401  (fixture reused)


def login(appmod, username, password):
    other = appmod.app.test_client()
    response = other.post("/login", data={"username": username, "password": password})
    assert response.status_code == 302 and "/login" not in response.headers["Location"], username
    return other


@pytest.fixture()
def chain(client):
    import app as appmod

    submit_review(client, upload(client, [SAMPLES / "D_chain_multi_branch.xlsx"]))
    for username, role, branches in [("zayed_mgr", "manager", ["Zayed"]), ("viewer1", "viewer", ["Downtown", "Nasr City"]), ("nobody", "viewer", [])]:
        client.post("/users", data={"action": "create", "username": username, "password": "password123", "role": role, "branches": branches})
    return appmod, client


def test_manager_only_sees_own_branch_everywhere(chain):
    appmod, _owner = chain
    manager = login(appmod, "zayed_mgr", "password123")
    for url in ["/", "/?branch=Downtown", "/reports/all", "/reports/all?branch=Downtown", "/profit-loss?branch=Downtown",
                "/abc", "/trends", "/financial", "/uploads", "/products/C1000", "/actions?branch=Downtown"]:
        page = manager.get(url).data.decode()
        if url.startswith(("/reports", "/uploads")):
            assert "<td>Zayed</td>" in page, url
        if not url.startswith("/actions"):  # transfers name the other branch on purpose
            assert "Downtown" not in page and "Nasr City" not in page, url
        assert "<option value=\"Downtown\"" not in page and 'value="all"' not in page, url
    exported = manager.get("/export/all?branch=Downtown").data.decode()
    branches = {line.split(",")[1] for line in exported.splitlines()[1:] if line}
    assert branches == {"Zayed"}
    # Transfers name the other branch (that is the point) but orders are only Zayed's.
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(manager.get("/purchase-orders.xlsx").data))
    assert all(name.startswith("Zayed") for name in workbook.sheetnames)


def test_manager_cannot_reach_admin_pages_or_upload_other_branches(chain):
    appmod, _owner = chain
    manager = login(appmod, "zayed_mgr", "password123")
    for url in ["/settings", "/users", "/data-management"]:
        assert manager.get(url).status_code == 403, url
    assert manager.post("/settings", data={"action": "reset"}).status_code == 403
    # The chain file holds three branches: a Zayed manager may not import it.
    response = manager.post("/upload", data={"files": [(SAMPLES.joinpath("D_chain_multi_branch.xlsx").open("rb"), "D.xlsx")]}, content_type="multipart/form-data")
    review = response.headers["Location"]
    rejected = submit_review(manager, review)
    assert "/upload/review/" in rejected.headers["Location"]
    # A single-branch file goes to the manager's own branch even without typing it.
    response = manager.post("/upload", data={"files": [(SAMPLES.joinpath("A_arabic_pharmacy_report.xlsx").open("rb"), "A.xlsx")]}, content_type="multipart/form-data")
    accepted = submit_review(manager, response.headers["Location"], {"branch": ""})
    assert "/imports/" in accepted.headers["Location"]
    with appmod.get_db() as db:
        assert {row[0] for row in db.execute("SELECT DISTINCT branch FROM items")} == {"Downtown", "Nasr City", "Zayed"}


def test_viewer_is_read_only_and_branchless_user_sees_nothing(chain):
    appmod, _owner = chain
    viewer = login(appmod, "viewer1", "password123")
    assert viewer.get("/upload").status_code == 403
    assert "Nasr City" in viewer.get("/?branch=Nasr City").data.decode()
    assert viewer.get("/reports/all?branch=Zayed").data.decode().count("<td>Zayed</td>") == 0
    nobody = login(appmod, "nobody", "password123")
    assert nobody.get("/reports/all").data.decode().count("item-link") == 0


def test_last_owner_is_protected(chain):
    appmod, owner = chain
    with appmod.get_db() as db:
        owner_id = db.execute("SELECT id FROM users WHERE username = 'admin'").fetchone()[0]
    owner.post("/users", data={"action": "update", "user_id": owner_id, "role": "viewer"})
    owner.post("/users", data={"action": "delete", "user_id": owner_id, "role": "owner"})
    with appmod.get_db() as db:
        assert db.execute("SELECT role FROM users WHERE id = ?", (owner_id,)).fetchone()[0] == "owner"
