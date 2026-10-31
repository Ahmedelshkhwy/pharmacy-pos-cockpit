"""End-to-end web flow: upload -> review -> analyze -> pages, on an isolated data folder."""
from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "sample_data"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(SAMPLES))

import generate_samples  # noqa: E402


@pytest.fixture()
def client(tmp_path, monkeypatch):
    generate_samples.main()
    import app as appmod

    data = tmp_path / "instance"
    for name, value in {
        "DATA_DIR": data,
        "UPLOAD_DIR": data / "uploads",
        "POS_RAW_DIR": data / "pos_raw",
        "POS_ANALYSIS_DIR": data / "pos_analysis",
        "IMPORTS_DIR": data / "imports",
        "MAPPING_PROFILES_PATH": data / "mapping_profiles.json",
        "DB_PATH": data / "pharmacy.db",
    }.items():
        monkeypatch.setattr(appmod, name, value)
    appmod.init_db()
    appmod.app.config["TESTING"] = True
    test_client = appmod.app.test_client()
    test_client.post("/login", data={"username": "admin", "password": "admin123"})
    return test_client


def submit_review(client, review_url: str, overrides: dict[str, str] | None = None):
    html = client.get(review_url).data.decode()
    form = {name: "1" for name in re.findall(r'name="(use-[^"]+)" value="1" checked', html)}
    for name, body in re.findall(r'<select name="(map-[^"]+)">(.*?)</select>', html, flags=re.S):
        selected = re.search(r'<option value="([^"]*)" selected', body)
        form[name] = selected.group(1) if selected else ""
    branch = re.search(r'name="branch"[^>]*value="([^"]*)"', html)
    form["branch"] = branch.group(1) if branch else ""
    form.update(overrides or {})
    return client.post(review_url, data=form)


def upload(client, paths, branch=""):
    data = {"branch": branch, "files": [(path.open("rb"), path.name) for path in paths]}
    response = client.post("/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 302 and "/upload/review/" in response.headers["Location"]
    return response.headers["Location"]


def test_upload_review_analyze_and_browse(client):
    review = upload(client, [SAMPLES / "B_sales_invoices.csv", SAMPLES / "B_stock_batches.csv"], branch="مصر الجديدة")
    assert "Quantity sold" in client.get(review).data.decode()
    done = submit_review(client, review)
    assert "/imports/" in done.headers["Location"]
    result = client.get(done.headers["Location"]).data.decode()
    assert "مصر الجديدة" in result

    for url in ["/", "/profit-loss", "/reports/expiry", "/reports/reorder", "/reports/loss", "/financial",
                "/trends", "/trends?branch=مصر الجديدة", "/settings", "/uploads", "/data-management", "/export/expiry",
                "/actions", "/abc", "/purchase-orders.xlsx"]:
        assert client.get(url).status_code == 200, url
    page = client.get("/profit-loss?branch=مصر الجديدة").data.decode()
    in_stock_file = [item for item in generate_samples.build_branch(23, "B") if item.stock > 0]  # cost comes from there
    for name in generate_samples.below_cost_names(in_stock_file):
        assert name in page  # the items sold below cost are named


def test_user_can_fix_a_misread_column_and_layout_is_remembered(client):
    path = SAMPLES / "D_chain_multi_branch.xlsx"
    review = upload(client, [path])
    # Pretend the store column was wrong: ignore it, everything goes to one branch.
    html = client.get(review).data.decode()
    store_select = re.search(r'<td><strong>Store</strong></td>\s*<td>\s*<select name="(map-[^"]+)"', html).group(1)
    done = submit_review(client, review, {store_select: "_ignore", "branch": "One Branch"})
    assert "/imports/" in done.headers["Location"]
    assert client.get("/profit-loss?branch=One Branch").status_code == 200

    second = client.get(upload(client, [path])).data.decode()
    assert "saved layout" in second or "شكل ملف محفوظ" in second


def test_period_is_asked_when_files_do_not_say(client, tmp_path):
    path = tmp_path / "summary.csv"
    path.write_text("Item,Sold,Stock,Cost,Price\nPanadol,10,5,30,40\n", encoding="utf-8")
    review = upload(client, [path], branch="Test")
    assert "required" in client.get(review).data.decode()
    assert "/upload/review/" in submit_review(client, review).headers["Location"]
    assert "/imports/" in submit_review(client, review, {"period_days": "30"}).headers["Location"]


def test_rejects_unsupported_file(client, tmp_path):
    bad = tmp_path / "notes.pdf"
    bad.write_bytes(b"%PDF")
    response = client.post("/upload", data={"files": [(bad.open("rb"), bad.name)]}, content_type="multipart/form-data")
    assert response.headers["Location"].endswith("/upload")


def snapshot_count(branch: str) -> int:
    import app as appmod

    with appmod.get_db() as db:
        return db.execute("SELECT COUNT(*) FROM snapshots WHERE branch = ?", (branch,)).fetchone()[0]


def test_every_upload_is_kept_as_history(client):
    branch = "المعادي"
    files = sorted((SAMPLES / "E_monthly_uploads").glob("*.xlsx"))
    for path in files:
        submit_review(client, upload(client, [path], branch=branch))
    assert snapshot_count(branch) == 3

    page = client.get(f"/trends?branch={branch}").data.decode()
    assert all(day in page for day in ["2026-08-31", "2026-09-30", "2026-10-30"])  # all three periods are listed
    assert "Losses went up" in page or "Losses went down" in page  # item-level movers between the last two
    history = client.get(f"/products/A1000?branch={branch}").data.decode()
    assert history.count("<td>2026-") == 3

    # Re-uploading the same period replaces that snapshot instead of duplicating it.
    submit_review(client, upload(client, [files[-1]], branch=branch))
    assert snapshot_count(branch) == 3
    assert client.get("/weekly-comparison").headers["Location"].endswith("/trends")


def test_settings_change_results_after_reanalysis(client):
    import app as appmod

    submit_review(client, upload(client, [SAMPLES / "A_arabic_pharmacy_report.xlsx"], branch="Maadi"))
    with appmod.get_db() as db:
        before = {row["sku"]: row["reorder_point"] for row in db.execute("SELECT sku, reorder_point FROM items")}

    form = {key: str(value) for key, value in appmod.AnalysisSettings().to_dict().items() if key != "lead_time_overrides"}
    form.update({"lead_time_days": "14", "correct_stockouts": "1", "override_kind": "category", "override_name": "Antibiotics", "override_days": "2"})
    client.post("/settings", data=form)
    settings = appmod.load_analysis_settings()
    assert settings.lead_time_days == 14 and settings.lead_time_overrides == [{"kind": "category", "name": "Antibiotics", "days": 2}]

    client.post("/settings", data={"action": "reanalyze"})
    with appmod.get_db() as db:
        rows = db.execute("SELECT sku, category, lead_time_days, reorder_point, sold FROM items").fetchall()
    assert {row["lead_time_days"] for row in rows if row["category"] == "Antibiotics"} == {2}
    assert {row["lead_time_days"] for row in rows if row["category"] != "Antibiotics"} == {14}
    selling = [row for row in rows if row["sold"] > 0 and row["category"] != "Antibiotics"]
    assert all(row["reorder_point"] > before[row["sku"]] for row in selling)  # longer lead time -> reorder earlier
    assert snapshot_count("Maadi") == 1  # same period re-analyzed, not duplicated


def test_action_center_and_purchase_order_excel(client):
    from io import BytesIO

    from openpyxl import load_workbook

    submit_review(client, upload(client, [SAMPLES / "D_chain_multi_branch.xlsx"]))
    page = client.get("/actions").data.decode()
    assert "wa.me" in page and "GSK Distribution" in page
    assert "Downtown" in page and "Nasr City" in page  # transfers name both ends

    response = client.get("/purchase-orders.xlsx?branch=Zayed")
    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames and all(name.startswith("Zayed") for name in workbook.sheetnames)
    sheet = workbook[workbook.sheetnames[0]]
    quantities = [row[5] for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]
    total_row = [row for row in sheet.iter_rows(values_only=True) if row[2] in ("Total", "الإجمالي")][0]
    assert quantities and total_row[5] == sum(quantities)
